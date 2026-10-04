import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Union
from pathlib import Path
from datetime import datetime
import pandas as pd

class ExcelGenerator:
    """
    Generates structured, production-grade Excel workbooks (.xlsx) from parsed voter records.
    Ensures Marathi Unicode fidelity and strictly prevents generating empty Excel files.
    """

    @classmethod
    def generate_voter_excel(
        cls,
        records: List[Dict[str, Any]],
        output_path: Union[str, Path],
        title: str = "Voter Data"
    ) -> Path:
        out_path = Path(output_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        if not records:
            raise ValueError(
                "Cannot generate Excel file: no records were extracted from the PDF. "
                "Generating empty Excel files is prohibited."
            )

        wb = openpyxl.Workbook()

        # ----------------- SHEET 1: VOTERS LIST (मतदार यादी) -----------------
        ws_data = wb.active
        ws_data.title = "Voters List (मतदार यादी)"

        # Expected Column Structure as requested by specification:
        # Sr No | Full Name Marathi | Surname Marathi | First Name Marathi | Father Name Marathi | Membership Number | Village Marathi | Mobile Number
        headers = [
            ("Sr No", "अ.क्र."),
            ("Full Name Marathi", "पूर्ण नाव मराठी"),
            ("Surname Marathi", "आडनाव मराठी"),
            ("First Name Marathi", "नाव मराठी"),
            ("Father Name Marathi", "वडिलांचे नाव मराठी"),
            ("Membership Number", "सभासद क्र."),
            ("Village Marathi", "गाव मराठी"),
            ("Mobile Number", "मोबाईल क्र."),
            ("EPIC / ID", "ओळखपत्र क्र."),
            ("House No", "घर क्र."),
            ("Age", "वय"),
            ("Gender", "लिंग"),
            ("Status", "स्थिती")
        ]

        # Style Definitions
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Dark Navy
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Arial", size=10)
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Write Header Row
        for col_idx, (en_h, mr_h) in enumerate(headers, 1):
            cell = ws_data.cell(row=1, column=col_idx, value=f"{en_h}\n({mr_h})")
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_thin

        ws_data.row_dimensions[1].height = 36

        # Write Data Rows
        for r_idx, rec in enumerate(records, start=2):
            sr_val = rec.get("serial_number") or (r_idx - 1)
            full_name = rec.get("full_name_mr") or rec.get("full_name_en") or ""
            surname = rec.get("surname") or ""
            first_name = rec.get("first_name") or ""
            father_name = rec.get("father_name") or rec.get("relative_name_mr") or ""
            mem_no = rec.get("membership_number") or rec.get("epic_number") or ""
            village = rec.get("village_mr") or rec.get("village") or ""
            mobile = rec.get("mobile_number") or ""
            epic = rec.get("epic_number") or rec.get("membership_number") or ""
            house = rec.get("house_number") or ""
            age = rec.get("age") or ""
            gender = rec.get("gender") or "Male"
            status = rec.get("validation_status") or rec.get("status") or "ACTIVE"

            row_data = [
                sr_val,
                full_name,
                surname,
                first_name,
                father_name,
                mem_no,
                village,
                mobile,
                epic,
                house,
                age,
                gender,
                status
            ]

            is_even = (r_idx % 2 == 0)
            for c_idx, val in enumerate(row_data, 1):
                cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
                cell.font = cell_font
                cell.border = border_thin
                if is_even:
                    cell.fill = zebra_fill

                # Center align identifiers, numbers, age, gender, status
                if c_idx in (1, 6, 8, 9, 10, 11, 12, 13):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_data.row_dimensions[r_idx].height = 22

        # Auto-adjust column widths
        for col in ws_data.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                # Handle newlines in header
                lines = val_str.split("\n")
                line_max = max(len(l) for l in lines) if lines else 0
                if line_max > max_len:
                    max_len = line_max
            ws_data.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # ----------------- SHEET 2: SUMMARY (सारांश) -----------------
        ws_sum = wb.create_sheet(title="Summary (सारांश)")
        ws_sum.column_dimensions["A"].width = 28
        ws_sum.column_dimensions["B"].width = 35

        total_records = len(records)
        valid_count = sum(1 for r in records if r.get("validation_status") == "VALID" or r.get("is_valid") is True)
        invalid_count = sum(1 for r in records if r.get("validation_status") == "INVALID" or r.get("is_valid") is False)
        duplicate_count = sum(1 for r in records if r.get("is_duplicate") or r.get("is_batch_duplicate"))

        sum_rows = [
            ("Report Title", title),
            ("Generated On", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Total Records Found", total_records),
            ("Valid Records (वैध नोंदी)", valid_count),
            ("Invalid Records (त्रुटीपूर्ण नोंदी)", invalid_count),
            ("Duplicate Records (दुबार नोंदी)", duplicate_count),
            ("Status", "SUCCESS" if total_records > 0 else "EMPTY")
        ]

        sum_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        for idx, (lbl, val) in enumerate(sum_rows, 1):
            c1 = ws_sum.cell(row=idx, column=1, value=lbl)
            c2 = ws_sum.cell(row=idx, column=2, value=val)
            c1.font = Font(name="Arial", bold=True)
            c2.font = Font(name="Arial")
            c1.border = border_thin
            c2.border = border_thin
            if idx == 1:
                c1.fill = sum_header_fill
                c1.font = Font(name="Arial", bold=True, color="FFFFFF")

        wb.save(str(out_path))
        return out_path

def generate_excel(
    records: List[Dict[str, Any]],
    output_path: Union[str, Path],
    title: str = "Electoral Roll Extraction"
) -> Path:
    """
    Pipeline entry point for generating Excel from extracted records.
    """
    return ExcelGenerator.generate_voter_excel(records=records, output_path=output_path, title=title)
