import pytest
from pathlib import Path
import openpyxl

from app.pdf_processing.detector import detect_pdf_type, inspect_pdf_metadata
from app.pdf_processing.marathi_normalizer import (
    to_arabic_num,
    normalize_marathi_text,
    clean_mobile_number,
    clean_membership_number,
    parse_marathi_name
)
from app.pdf_processing.extractor import extract_table_from_pdf
from app.pdf_processing.table_parser import parse_records
from app.pdf_processing.mapper import map_columns
from app.pdf_processing.validator import validate_records
from app.pdf_processing.excel_generator import generate_excel
from app.pdf_processing.pipeline import process_pdf

SAMPLE_PDF = Path("uploads/Sakharele_Marathi_Voter_Table.pdf")
SCANNED_PDF = Path("uploads/58.pdf")

def test_marathi_normalizer():
    # Marathi numerals to Arabic
    assert to_arabic_num("१२३४५") == "12345"
    assert to_arabic_num("०९८७६") == "09876"

    # Non-breaking space and Unicode hyphens
    raw_name = "पाटील\xa0सचिन\xa0आनंदराव"
    assert normalize_marathi_text(raw_name) == "पाटील सचिन आनंदराव"
    assert clean_membership_number("B‐1") == "B-1"

    # Mobile cleaning
    assert clean_mobile_number("+91 91724 74077") == "9172474077"
    assert clean_mobile_number("०९८२२३३४४५५") == "9822334455"

    # Name parts
    s, f, fa = parse_marathi_name("पाटील सचिन आनंदराव")
    assert s == "पाटील"
    assert f == "सचिन"
    assert fa == "आनंदराव"

def test_detector():
    if SAMPLE_PDF.exists():
        assert detect_pdf_type(SAMPLE_PDF) == "TEXT"
        meta = inspect_pdf_metadata(SAMPLE_PDF)
        assert meta["has_selectable_text"] is True
        assert meta["total_pages"] >= 1

    if SCANNED_PDF.exists():
        assert detect_pdf_type(SCANNED_PDF) == "SCANNED"

def test_table_parser():
    if not SAMPLE_PDF.exists():
        pytest.skip("Sample PDF not found")

    raw_data = extract_table_from_pdf(SAMPLE_PDF)
    records = parse_records(raw_data)
    assert len(records) == 5

    rec1 = records[0]
    assert rec1["serial_number"] == 1
    assert "सचिन" in rec1["full_name_mr"]
    assert rec1["surname"] == "पाटील"
    assert rec1["first_name"] == "सचिन"
    assert rec1["father_name"] == "आनंदराव"
    assert rec1["membership_number"] == "B-1"
    assert rec1["village_mr"] == "साखराळे"
    assert rec1["mobile_number"] == "9172474077"

def test_mapper_and_validator():
    raw_row = {
        "Sr No": "1",
        "Full Name Marathi": "पाटील सचिन आनंदराव",
        "Surname Marathi": "पाटील",
        "First Name Marathi": "सचिन",
        "Father Name Marathi": "आनंदराव",
        "Membership Number": "B‐1",
        "Village Marathi": "साखराळे",
        "Mobile Number": "9172474077"
    }
    mapped = map_columns([raw_row])[0]
    assert mapped["serial_number"] == 1
    assert mapped["full_name_mr"] == "पाटील सचिन आनंदराव"
    assert mapped["membership_number"] == "B-1"
    assert mapped["epic_number"] == "B-1"
    assert mapped["mobile_number"] == "9172474077"

    validated = validate_records([mapped])[0]
    assert validated["is_valid"] is True
    assert validated["validation_status"] == "VALID"

def test_excel_generator_non_empty(tmp_path):
    records = [
        {
            "serial_number": 1,
            "full_name_mr": "पाटील सचिन आनंदराव",
            "surname": "पाटील",
            "first_name": "सचिन",
            "father_name": "आनंदराव",
            "membership_number": "B-1",
            "village_mr": "साखराळे",
            "mobile_number": "9172474077",
            "validation_status": "VALID"
        }
    ]
    out_xlsx = tmp_path / "test_voters.xlsx"
    res_path = generate_excel(records, out_xlsx)
    assert res_path.exists()
    assert res_path.stat().st_size > 2000

    wb = openpyxl.load_workbook(res_path)
    assert "Voters List (मतदार यादी)" in wb.sheetnames
    ws = wb["Voters List (मतदार यादी)"]
    assert ws.max_row == 2
    assert "Full Name Marathi" in str(ws.cell(1, 2).value)
    assert ws.cell(2, 2).value == "पाटील सचिन आनंदराव"
    assert ws.cell(2, 6).value == "B-1"

def test_excel_generator_empty_fails(tmp_path):
    out_xlsx = tmp_path / "empty.xlsx"
    with pytest.raises(ValueError) as exc:
        generate_excel([], out_xlsx)
    assert "no records were extracted" in str(exc.value)

def test_pipeline_end_to_end(tmp_path):
    if not SAMPLE_PDF.exists():
        pytest.skip("Sample PDF not found")

    out_xlsx = tmp_path / "pipeline_output.xlsx"
    result = process_pdf(SAMPLE_PDF, output_excel_path=out_xlsx)
    assert result["status"] == "COMPLETED"
    assert result["records_found"] == 5
    assert result["valid_records"] == 5
    assert result["invalid_records"] == 0
    assert result["duplicate_records"] == 0
    assert Path(result["excel_file"]).exists()

def test_pipeline_zero_records_fails():
    if not SCANNED_PDF.exists():
        pytest.skip("Scanned PDF not found")

    with pytest.raises(ValueError) as exc:
        process_pdf(SCANNED_PDF)
    assert "0 records found" in str(exc.value)
